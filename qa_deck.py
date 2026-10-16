"""QA deck.pptx: estimated text overflow, shapes outside slide, leftover placeholders."""
import math
from pptx import Presentation
from pptx.util import Emu

PATH = r"C:/Users/Prekzursil/.zcode/workspace/default/tss/course-ambiegen/deck.pptx"
SW, SH = 13.333, 7.5
# rough advance width factor (fraction of font size) per font
WFAC = {"Segoe UI": 0.50, "Consolas": 0.55}

prs = Presentation(PATH)
issues = []

def walk(shapes, sidx):
    for sp in shapes:
        if sp.shape_type == 6:
            walk(sp.shapes, sidx)
            continue
        x = Emu(sp.left).inches if sp.left is not None else 0
        y = Emu(sp.top).inches if sp.top is not None else 0
        w = Emu(sp.width).inches if sp.width is not None else 0
        h = Emu(sp.height).inches if sp.height is not None else 0
        if x < -0.01 or y < -0.01 or x + w > SW + 0.01 or y + h > SH + 0.01:
            issues.append(f"slide {sidx}: shape '{sp.shape_type}' out of slide ({x:.2f},{y:.2f},{w:.2f},{h:.2f})")
        if not getattr(sp, "has_text_frame", False) or not sp.has_text_frame:
            continue
        tf = sp.text_frame
        if not tf.text.strip():
            continue
        est_h = 0.0
        text_all = []
        for p in tf.paragraphs:
            txt = "".join(r.text for r in p.runs)
            text_all.append(txt)
            if not txt:
                est_h += 0.14
                continue
            size = max((r.font.size.pt if r.font.size else 14) for r in p.runs)
            font = next((r.font.name for r in p.runs if r.font.name), "Segoe UI")
            fac = WFAC.get(font, 0.5)
            cpl = max(int(w / (size * fac / 72)), 1) if w > 0 else 1
            lines = max(math.ceil(len(txt) / cpl), txt.count("\n") + 1)
            ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
            est_h += lines * size * 1.22 * ls / 72
            if p.space_after:
                est_h += p.space_after.pt / 72
        joined = " ".join(text_all).lower()
        for bad in ("todo", "lorem", "xxx", "[insert", "click to add"):
            if bad in joined:
                issues.append(f"slide {sidx}: placeholder text '{bad}' in {text_all[:1]}")
        if est_h > h * 1.18 + 0.06:
            issues.append(f"slide {sidx}: possible overflow ({est_h:.2f}in est vs {h:.2f}in box) text={text_all[:1]}")

for i, sl in enumerate(prs.slides, 1):
    walk(sl.shapes, i)

print(f"slides: {len(prs.slides._sldIdLst)}")
if issues:
    print(f"{len(issues)} issue(s):")
    for it in issues:
        print(" -", it)
else:
    print("no issues found")
